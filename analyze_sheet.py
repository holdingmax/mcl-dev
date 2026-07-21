import openpyxl, re, json
from datetime import datetime, timedelta

wb = openpyxl.load_workbook(r'C:\Users\dgarc\Downloads\EXPRESS LINE  BOOKING.xlsx', data_only=True)
ws = wb['May 2026 ']

def parse_date(raw):
    if raw is None or raw == '': return ''
    if isinstance(raw, (int, float)):
        d = datetime(1899, 12, 30) + timedelta(days=raw)
        if raw % 1 != 0:
            return d.strftime('%Y-%m-%dT%H:%M:00')
        return d.strftime('%Y-%m-%d')
    if isinstance(raw, datetime):
        if raw.hour or raw.minute:
            return raw.strftime('%Y-%m-%dT%H:%M:00')
        return raw.strftime('%Y-%m-%d')
    s = str(raw).replace('\r', ' ').replace('\n', ' ').replace('\t', ' ')
    s = re.sub(r'\s+', ' ', s).strip()
    if not s or re.match(r'^[a-z]', s, re.I): return s
    m = re.search(r'(\d{1,2})[/\-](\d{1,2})(?:[/\-](\d{2,4}))?', s)
    if m:
        mm = m.group(1).zfill(2)
        dd = m.group(2).zfill(2)
        yr = m.group(3)
        if yr:
            year = '20' + yr if len(yr) == 2 else yr
        else:
            year = str(datetime.now().year)
        return f'{year}-{mm}-{dd}'
    return s

def is_yes(val):
    v = str(val or '').lower().strip()
    return v in ('y', 'yes', 'si', 'true', '1')

def norm(s): return str(s or '').strip().upper()

COL = {
    'inbond': 0, 'doc_status': 1, 'flight': 2, 'awb1': 3, 'awb2': 4,
    'uld': 5, 'pcs': 6, 'cw': 7, 'pickup': 8, 'ready': 9, 'lfd': 10,
    'deliver_to': 11, 'dest': 12, 'cutoff': 13, 'etd': 14,
    'hold_bg': 15, 'hold_box': 16, 'battery': 17, 'on_hand': 18,
    'segregation': 19, 'dropped_off': 20, 'vuelo_num': 21
}

def cell(r, key):
    return str(ws.cell(r, COL[key]+1).value or '').strip()

def cell_raw(r, key):
    return ws.cell(r, COL[key]+1).value

rows = []
cont_holds = []
last_awb_idx = -1

for r in range(2, ws.max_row + 1):
    awb1 = cell(r, 'awb1')
    awb2 = cell(r, 'awb2')
    hold_bg = cell(r, 'hold_bg')
    hold_box = cell(r, 'hold_box')

    if not awb1 and not awb2:
        if hold_bg and last_awb_idx >= 0:
            hbl = hold_bg.lower().strip()
            if hbl and hbl not in ('no hold', 'no', 'n/a'):
                cont_holds.append({'parent': last_awb_idx, 'bag': hold_bg, 'box': hold_box})
        continue

    awb_check = (awb1 or awb2).strip()
    if not re.search(r'\d{3}-\d+', awb_check):
        continue

    last_awb_idx = len(rows)

    flight = norm(cell(r, 'flight'))
    pcs_raw = cell_raw(r, 'pcs')
    cw_raw = cell_raw(r, 'cw')

    bg = hold_bg.strip()
    bg_lower = bg.lower()
    is_no = not bg or bg_lower in ('no hold', 'no', 'n/a')
    is_pend = bg_lower in ('pending', 'pendiente')

    hold_list = []
    if not is_no and not is_pend and bg:
        hold_list.append({'bag': bg, 'box': hold_box})

    h_status = 'PENDING' if is_pend else ('HOLD' if not is_no and not is_pend and bg else 'NO HOLD')

    rows.append({
        'row': r,
        'flight': flight,
        'awb1': awb1,
        'awb2': awb2,
        'uld': norm(cell(r, 'uld')) or 'LOOSE',
        'pcs': int(float(pcs_raw)) if pcs_raw else 0,
        'cw': float(cw_raw) if cw_raw else 0,
        'pickup': cell(r, 'pickup'),
        'ready': parse_date(cell_raw(r, 'ready')),
        'lfd': parse_date(cell_raw(r, 'lfd')),
        'deliver_to': cell(r, 'deliver_to'),
        'dest': norm(cell(r, 'dest')),
        'cutoff': parse_date(cell_raw(r, 'cutoff')),
        'etd': parse_date(cell_raw(r, 'etd')),
        'hold_bg': hold_bg,
        'hold_box': hold_box,
        'battery': is_yes(cell(r, 'battery')),
        'docs': 'DOCS' in cell(r, 'doc_status').upper(),
        'on_hand_raw': cell(r, 'on_hand'),
        'segregation': cell(r, 'segregation'),
        'dropped_raw': cell(r, 'dropped_off'),
        'holds': hold_list,
        'hold_status': h_status,
    })

for ch in cont_holds:
    if 0 <= ch['parent'] < len(rows):
        rows[ch['parent']]['holds'].append({'bag': ch['bag'], 'box': ch['box']})
        if rows[ch['parent']]['hold_status'] == 'NO HOLD':
            rows[ch['parent']]['hold_status'] = 'HOLD'

print(f'Total valid AWB rows: {len(rows)}')

# Group by flight
flights = {}
for row in rows:
    f = row['flight']
    if f not in flights: flights[f] = []
    flights[f].append(row)

print(f'Unique flights in sheet: {len(flights)}')
for f in sorted(flights.keys()):
    awbs = flights[f]
    dests = set(a['dest'] for a in awbs)
    etds = set(a['etd'] for a in awbs)
    print(f'  {f}: {len(awbs)} AWBs, dest={dests}, etd={etds}')

# FLIGHT NORMALIZATION ISSUES
print('\n=== FLIGHT NAME NORMALIZATION ===')
flight_variants = {}
for f in flights:
    base = re.sub(r'[\s\-]', '', f)
    if base not in flight_variants: flight_variants[base] = set()
    flight_variants[base].add(f)
for base, variants in sorted(flight_variants.items()):
    if len(variants) > 1:
        total = sum(len(flights[v]) for v in variants)
        print(f'  CONFLICT: {sorted(variants)} ({total} total AWBs) -> creates DUPLICATE flights!')

# MISSING 2ND LEG
no_2nd = [r for r in rows if not r['awb2']]
print(f'\n=== MISSING 2ND LEG ({len(no_2nd)} rows - SKIPPED by agent) ===')
for r in no_2nd:
    print(f'  R{r["row"]}: {r["awb1"]} flight={r["flight"]} dest={r["dest"]} pcs={r["pcs"]}')

# MISSING ETD
no_etd = [r for r in rows if not r['etd'] and r['awb2']]
print(f'\n=== MISSING ETD with 2nd leg ({len(no_etd)} rows) ===')
for r in no_etd:
    print(f'  R{r["row"]}: {r["awb1"]} flight={r["flight"]}')

# DATE TIME LOSS
print('\n=== TIME LOSS IN DATE PARSING ===')
time_lost = 0
for r in rows:
    for field, key in [('ready', 'ready'), ('cutoff', 'cutoff'), ('etd', 'etd')]:
        raw = str(cell_raw(r['row'], key) or '')
        parsed = r[field]
        if re.search(r'\d+:\d+', raw) and parsed and 'T' not in parsed:
            time_lost += 1
            if time_lost <= 10:
                print(f'  R{r["row"]} {field}: "{raw[:50]}" -> "{parsed}"')
print(f'  Total time-loss cases: {time_lost}')

# DEST ISSUES
print('\n=== DESTINATION NORMALIZATION ===')
all_dests = set(r['dest'] for r in rows if r['dest'])
for d in sorted(all_dests):
    count = sum(1 for r in rows if r['dest'] == d)
    issues = []
    if d != d.strip(): issues.append('trailing space')
    if ' ' in d.strip(): issues.append('contains space')
    note = f' *** {", ".join(issues)}' if issues else ''
    print(f'  {repr(d)}: {count} AWBs{note}')

# HOLDS
print('\n=== HOLDS ANALYSIS ===')
hold_rows = [r for r in rows if r['hold_status'] == 'HOLD']
pend_rows = [r for r in rows if r['hold_status'] == 'PENDING']
no_hold_rows = [r for r in rows if r['hold_status'] == 'NO HOLD']
print(f'  NO HOLD: {len(no_hold_rows)}, PENDING: {len(pend_rows)}, HOLD: {len(hold_rows)}')
for r in hold_rows:
    print(f'  R{r["row"]}: {r["awb1"]} holds={len(r["holds"])} bag={r["hold_bg"]}')
multi = [r for r in rows if len(r['holds']) > 1]
if multi:
    print(f'  Multi-hold (continuation rows merged): {len(multi)}')
    for r in multi:
        print(f'    R{r["row"]}: {r["awb1"]} -> {len(r["holds"])} holds')

# SEGREGATION
seg_rows = [r for r in rows if r['segregation']]
print(f'\n=== SEGREGATION (UNMAPPED, {len(seg_rows)} rows) ===')
segs = {}
for r in seg_rows:
    s = r['segregation']
    if s not in segs: segs[s] = 0
    segs[s] += 1
for s in sorted(segs.keys()):
    print(f'  "{s}": {segs[s]} AWBs')

# SPECIAL CASES
print('\n=== SPECIAL CASES ===')
for r in rows:
    inbond = str(cell_raw(r['row'], 'inbond') or '')
    if 'CBP' in inbond.upper() or r['awb1'] == '160-11964945':
        print(f'  R{r["row"]}: CBP hold - {r["awb1"]}, flight={r["flight"]}, holds={r["holds"]}')
    if r['pickup'] and ('pick up by' in r['pickup'].lower() or 'elc-bst' in r['pickup'].lower()):
        print(f'  R{r["row"]}: pickup by ELC-BST - {r["awb1"]}, ready={r["ready"]}')

# AGENT ELIGIBILITY
print('\n=== AGENT CREATION ELIGIBILITY ===')
eligible = []
skipped = []
for r in rows:
    has_flight = bool(r['flight'])
    has_etd = bool(r['etd'])
    has_2nd = bool(r['awb2'])
    if has_flight and has_etd and has_2nd:
        eligible.append(r)
    else:
        reasons = []
        if not has_flight: reasons.append('no flight')
        if not has_etd: reasons.append('no ETD')
        if not has_2nd: reasons.append('no 2nd leg')
        skipped.append((r, reasons))

print(f'  Eligible for creation: {len(eligible)}')
print(f'  Skipped (incomplete): {len(skipped)}')
for r, reasons in skipped:
    print(f'    R{r["row"]}: {r["awb1"]} - {", ".join(reasons)}')

# IDEAL FLIGHT TABLE
print('\n=== IDEAL FLIGHT TABLE ===')
ideal_flights = {}
for r in eligible:
    fn = r['flight']
    etd_date = r['etd'][:10] if r['etd'] else 'NONE'
    key = f'{fn}|{etd_date}'
    if key not in ideal_flights: ideal_flights[key] = []
    ideal_flights[key].append(r)

for key in sorted(ideal_flights.keys()):
    awbs = ideal_flights[key]
    fn, etd = key.split('|')
    dests = set(a['dest'] for a in awbs)
    print(f'\n  Flight: {fn} | ETD: {etd} | {len(awbs)} AWBs | Dest: {dests}')
    for a in awbs:
        batt = 'Y' if a['battery'] else 'N'
        hs = a['hold_status'][0]
        print(f'    R{a["row"]}: {a["awb1"]:>15} -> {a["awb2"]:>15} | {a["pcs"]:>3}pcs {a["cw"]:>7.1f}kg | {a["uld"]:<5} | H:{hs} B:{batt} | {a["dest"]}')

# DATE ANOMALIES
print('\n=== DATE ANOMALIES ===')
for r in rows:
    for f in ['etd', 'cutoff', 'ready', 'lfd']:
        v = r.get(f, '')
        if v and re.match(r'\d{4}', v) and not v.startswith('202'):
            print(f'  R{r["row"]} {f}={v} (BAD YEAR)')

# BATTERY distribution
batt_y = sum(1 for r in rows if r['battery'])
batt_n = sum(1 for r in rows if not r['battery'])
print(f'\n=== BATTERY: Y={batt_y}, N={batt_n} ===')

# PICKUP FROM distribution
pickups = {}
for r in rows:
    p = r['pickup'] if r['pickup'] else '(empty)'
    if p not in pickups: pickups[p] = 0
    pickups[p] += 1
print('\n=== PICKUP FROM ===')
for p in sorted(pickups.keys()):
    print(f'  {p}: {pickups[p]}')

# Summary count
print(f'\n=== SUMMARY ===')
print(f'Total AWB rows in sheet: {len(rows)}')
print(f'Eligible for agent creation: {len(eligible)}')
print(f'Skipped (incomplete): {len(skipped)}')
print(f'Unique flights (as-is): {len(flights)}')
print(f'Unique flights (normalized): {len(flight_variants)}')
print(f'With holds: {len(hold_rows)}, Pending holds: {len(pend_rows)}')
print(f'With battery: {batt_y}')
print(f'With segregation: {len(seg_rows)}')
print(f'Time-loss parsing: {time_lost} cases')
